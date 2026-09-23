import feedparser, json, os, requests, datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

FEEDS = ["https://dronedj.com/feed", "https://dronelife.com/feed", "https://dronexl.co/feed", "https://www.suasnews.com/feed"]
NEWSAPI_QUERY = "drone"
NEWSAPI_KEY = os.environ.get("NEWSAPI_KEY")
DIR = os.path.dirname(__file__)
HISTORY_FILE = os.path.join(DIR, "sent_history.json")
SHEET_FILE = os.path.join(DIR, "drone_news_log.xlsx")
OLLAMA_API = "http://localhost:11434/api/generate"
OLLAMA_MODEL = os.environ.get("OLLAMA_MODEL", "llama3.1")

def load_history():
    if not os.path.exists(HISTORY_FILE): return set()
    return set(json.load(open(HISTORY_FILE)))

def save_history(h):
    json.dump(list(h), open(HISTORY_FILE, "w"))

def fetch_rss(history):
    items = []
    for url in FEEDS:
        for e in feedparser.parse(url).entries:
            if e.link not in history:
                items.append({"title": e.title, "link": e.link, "summary": e.get("summary", e.title)})
    return items

def fetch_newsapi(history):
    if not NEWSAPI_KEY:
        return []
    r = requests.get("https://newsapi.org/v2/everything",
        params={"q": NEWSAPI_QUERY, "sortBy": "publishedAt", "language": "en", "pageSize": 20},
        headers={"X-Api-Key": NEWSAPI_KEY})
    r.raise_for_status()
    items = []
    for a in r.json().get("articles", []):
        if a["url"] not in history:
            items.append({"title": a["title"], "link": a["url"], "summary": a.get("description") or a.get("content") or a["title"]})
    return items

def fetch_candidates(history):
    return fetch_rss(history) + fetch_newsapi(history)

def summarize(item):
    prompt = f"Summarize this drone industry news article in 2-3 sentences for a busy reader. Title: {item['title']}\n\nContent: {item['summary']}"
    r = requests.post(OLLAMA_API, json={"model": OLLAMA_MODEL, "prompt": prompt, "stream": False})
    r.raise_for_status()
    return r.json()["response"].strip()

def append_row(date, title, summary, link):
    if os.path.exists(SHEET_FILE):
        wb = load_workbook(SHEET_FILE)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Drone News"
        headers = ["Date", "Title", "Summary", "Link"]
        for col, h in enumerate(headers, start=1):
            c = ws.cell(row=1, column=col, value=h)
            c.font = Font(name="Arial", bold=True)
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 40
        ws.column_dimensions["C"].width = 70
        ws.column_dimensions["D"].width = 50
    row = ws.max_row + 1
    values = [date, title, summary, link]
    for col, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=col, value=v)
        c.font = Font(name="Arial")
    wb.save(SHEET_FILE)

def main():
    history = load_history()
    candidates = fetch_candidates(history)
    if not candidates:
        print("No new articles today")
        return
    pick = candidates[0]
    summary = summarize(pick)
    append_row(datetime.date.today().isoformat(), pick["title"], summary, pick["link"])
    history.add(pick["link"])
    save_history(history)
    print("Logged:", pick["title"])

if __name__ == "__main__":
    main()
